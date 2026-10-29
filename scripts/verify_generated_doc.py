import docx
import zipfile

doc = docx.Document('WealthOS_Major_Project_Phase1_Documentation.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

bad_tokens = ['\ufffd', '[Guide Name]', '[HOD Name]', '[Project Coordinator]', 'PASTE SCREENSHOT', 'Sentinel']
found = 0
for i, p in enumerate(doc.paragraphs):
    for bt in bad_tokens:
        if bt.lower() in p.text.lower():
            print(f"Paragraph {i} contains '{bt}': {p.text[:70]}")
            found += 1

for t_idx, tbl in enumerate(doc.tables):
    for r_idx, row in enumerate(tbl.rows):
        for c_idx, cell in enumerate(row.cells):
            for bt in bad_tokens:
                if bt.lower() in cell.text.lower():
                    print(f"Table {t_idx} R{r_idx}C{c_idx} contains '{bt}': {cell.text[:60]}")
                    found += 1

if found == 0:
    print("VERIFIED: Zero placeholders, zero encoding artifacts, and zero leftover Sentinel references!")

with zipfile.ZipFile('WealthOS_Major_Project_Phase1_Documentation.docx') as z:
    media = [f for f in z.namelist() if f.startswith('word/media/')]
    print(f"Total embedded media files: {len(media)}")
    for m in media:
        print(f"  {m}")
