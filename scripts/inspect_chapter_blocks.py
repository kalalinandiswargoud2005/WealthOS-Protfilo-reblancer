import docx

doc = docx.Document('SentinelAI_Documentation (1).docx')

chapters = [
    ("COVER", 0, 25),
    ("CERTIFICATE", 25, 50),
    ("DECLARATION", 50, 65),
    ("ACKNOWLEDGEMENT", 65, 90),
    ("ABSTRACT", 90, 97),
    ("INDICES HEADINGS", 97, 103),
    ("CHAPTER 1", 103, 132),
    ("CHAPTER 2", 132, 176),
    ("CHAPTER 3", 176, 291),
    ("CHAPTER 4", 291, 345),
    ("CHAPTER 5", 345, 416),
    ("CHAPTER 6", 416, 449),
    ("REFERENCES", 449, len(doc.paragraphs))
]

for name, start, end in chapters:
    print(f"\n=================== {name} (P{start} to P{end}) ===================")
    for i in range(start, min(end, len(doc.paragraphs))):
        p = doc.paragraphs[i]
        txt = p.text.strip()
        if txt:
            print(f"P{i:3d} [{p.style.name}]: {txt[:90]}")
