import docx

doc = docx.Document('SentinelAI_Documentation (1).docx')

print("=== SECTIONS & HEADERS/FOOTERS ===")
for i, s in enumerate(doc.sections):
    h_text = ' '.join([p.text.strip() for p in s.header.paragraphs if p.text.strip()])
    f_text = ' '.join([p.text.strip() for p in s.footer.paragraphs if p.text.strip()])
    print(f"Section {i}:")
    print(f"  Header: '{h_text}'")
    print(f"  Footer: '{f_text}'")

print("\n=== PARAGRAPH STYLES USED ===")
styles_used = set(p.style.name for p in doc.paragraphs)
for st in sorted(styles_used):
    count = sum(1 for p in doc.paragraphs if p.style.name == st)
    print(f"  Style '{st}': {count} paragraphs")
