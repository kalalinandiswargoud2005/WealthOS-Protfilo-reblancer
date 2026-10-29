"""
build_wealthos_master_document.py
Master execution script that generates the complete, professional, publication-grade
Major Project Phase 1 Documentation for WealthOS.
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor

from scripts.build_complete_documentation import add_header_footer
from scripts.doc_content_prelims import add_preliminary_pages
from scripts.doc_content_ch1_ch2 import add_chapter_1, add_chapter_2
from scripts.doc_content_ch3 import add_chapter_3
from scripts.doc_content_ch4 import add_chapter_4
from scripts.doc_content_ch5_ch6 import add_chapter_5_and_6

def build_master_documentation():
    print("==================================================")
    print("STARTING WEALTHOS MASTER DOCUMENTATION GENERATION")
    print("==================================================")
    
    doc = Document()
    
    # Configure sections, margins, headers, and footers
    add_header_footer(doc)
    
    # 1. Preliminary Pages
    print("-> Adding Preliminary Pages (Cover, Certificate, Declaration, Acknowledgement, Abstract, Indices)...")
    add_preliminary_pages(doc)
    
    # 2. Chapter 1 & Chapter 2
    print("-> Adding Chapter 1 (Introduction) & Chapter 2 (System Analysis)...")
    add_chapter_1(doc)
    add_chapter_2(doc)
    
    # 3. Chapter 3
    print("-> Adding Chapter 3 (Architectural Design & 12 Architectural/Empirical Figures)...")
    add_chapter_3(doc)
    
    # 4. Chapter 4
    print("-> Adding Chapter 4 (Implementation, Code Snippets, API Routes & Test Cases)...")
    add_chapter_4(doc)
    
    # 5. Chapter 5 & Chapter 6 & References
    print("-> Adding Chapter 5 (Results & 9 Screenshots), Chapter 6 (Conclusion & Future Scope) & References...")
    add_chapter_5_and_6(doc)
    
    output_target1 = "WealthOS_Major_Project_Phase1_Documentation.docx"
    output_target2 = "WealthOS_Documentation.docx"
    
    print(f"-> Saving output to {output_target1}...")
    doc.save(output_target1)
    
    print(f"-> Saving backup copy to {output_target2}...")
    doc.save(output_target2)
    
    sz1 = os.path.getsize(output_target1)
    print("==================================================")
    print(f"SUCCESS: Generated {output_target1} ({sz1:,} bytes)")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    print("==================================================")

if __name__ == "__main__":
    build_master_documentation()
