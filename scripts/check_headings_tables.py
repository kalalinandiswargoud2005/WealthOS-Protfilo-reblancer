import docx

doc = docx.Document('WealthOS_Major_Project_Phase1_Documentation.docx')

print("=== VERIFYING HEADINGS IN GENERATED DOC ===")
for i, p in enumerate(doc.paragraphs):
    if p.runs and p.runs[0].bold and len(p.text.strip()) < 80:
        if any(p.text.strip().startswith(x) for x in ['CHAPTER', '1.', '2.', '3.', '4.', '5.', '6.', 'Algorithm', 'CERTIFICATE', 'DECLARATION', 'ACKNOWLEDGEMENT', 'ABSTRACT', 'INDEX', 'LIST OF']):
            print(f"P{i:3d}: {p.text.strip()[:75]}")

print("\n=== VERIFYING TABLES ===")
for t_idx, tbl in enumerate(doc.tables):
    rows = len(tbl.rows)
    cols = len(tbl.columns)
    c0 = tbl.rows[0].cells[0].text.strip().replace('\n', ' ')
    c1 = tbl.rows[0].cells[1].text.strip().replace('\n', ' ') if cols > 1 else ''
    print(f"Table {t_idx:2d} ({rows:2d}x{cols}): {c0[:35]} | {c1[:35]}")
