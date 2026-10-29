import docx

doc = docx.Document('SentinelAI_Documentation (1).docx')

print("--- AUDITING INDEX (TABLE 0) VS DOCUMENT HEADINGS ---")
index_entries = [row.cells[0].text.strip().replace('\n', ' ') for row in doc.tables[0].rows[1:]]
doc_headings = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith('Heading') or p.style.name.startswith('Tmpl Heading')]

print(f"Total Index entries: {len(index_entries)}")
print(f"Total Doc headings: {len(doc_headings)}")

# Check discrepancies
diagram_entries = [(row.cells[0].text.strip(), row.cells[1].text.strip()) for row in doc.tables[1].rows[1:]]
print("\n--- LIST OF DIAGRAMS (TABLE 1) ---")
for num, title in diagram_entries:
    print(f"  {num}: {title}")

# Check figures mentioned in document
fig_mentions = []
for p in doc.paragraphs:
    if p.text.strip().startswith('Figure'):
        fig_mentions.append(p.text.strip())

print("\n--- FIGURE CAPTIONS IN TEXT ---")
for fm in fig_mentions:
    print(f"  {fm}")

# Check screenshots
screenshot_entries = [(row.cells[0].text.strip(), row.cells[1].text.strip()) for row in doc.tables[2].rows[1:]]
print("\n--- LIST OF SCREENSHOTS (TABLE 2) ---")
for num, title in screenshot_entries:
    print(f"  {num}: {title}")
