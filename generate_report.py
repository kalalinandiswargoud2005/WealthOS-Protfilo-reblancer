import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = docx.Document()
    
    # Title formatting
    title_style = doc.styles['Title']
    title_font = title_style.font
    title_font.name = 'Arial'
    title_font.size = Pt(24)
    title_font.color.rgb = RGBColor(0, 51, 102)

    heading1_style = doc.styles['Heading 1']
    heading1_font = heading1_style.font
    heading1_font.name = 'Arial'
    heading1_font.size = Pt(16)
    heading1_font.color.rgb = RGBColor(0, 51, 102)
    
    heading2_style = doc.styles['Heading 2']
    heading2_font = heading2_style.font
    heading2_font.name = 'Arial'
    heading2_font.size = Pt(14)
    heading2_font.color.rgb = RGBColor(51, 102, 153)

    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Arial'
    normal_font.size = Pt(11)

    # Student Info
    doc.add_paragraph('Name: G Vaishnav Kumar')
    doc.add_paragraph('Roll No: 2311CS040053')
    doc.add_paragraph('')

    # Main Title
    title = doc.add_paragraph('SentinelAI', style='Title')
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = doc.add_paragraph('Applied Agentic AI Assignment Report\nProfessional Report')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.bold = True
    
    doc.add_paragraph('')
    categories = doc.add_paragraph('LLM Workflow  |  Prompt Chaining  |  Agentic AI  |  RAG')
    categories.alignment = WD_ALIGN_PARAGRAPH.CENTER
    categories.runs[0].font.italic = True
    categories.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    
    doc.add_paragraph('')

    # Abstract
    doc.add_heading('Abstract', level=1)
    doc.add_paragraph('SentinelAI is an AI-powered cybersecurity monitoring and Security Operations Center (SOC) platform that unifies Large Language Models, Agentic AI, Prompt Chaining and Retrieval-Augmented Generation (RAG) within a single web application. The platform continuously observes cyber events, correlates security incidents, profiles attacker behaviour and maps observed threats to the MITRE ATT&CK framework. Through an integrated AI Security Copilot, SentinelAI generates investigation reports, explains attack patterns and assists security analysts with both general conversation and deep, evidence-grounded cybersecurity investigation, thereby accelerating incident response and reducing analyst workload.')

    # Q1
    doc.add_heading('Q1. LLM Workflow', level=1)
    doc.add_paragraph("SentinelAI's AI Security Copilot allows an analyst to ask natural-language questions directly inside the SOC dashboard. Every query is submitted from the React frontend to the FastAPI backend, which first assembles a Context Builder — pulling in the relevant live incidents, honeypot events, prior investigation notes and threat-intelligence records associated with the request. This structured context, together with the user's question, is packaged into a prompt and sent to the Groq LLM. The model returns security recommendations, plain-language explanations of attacker techniques, or general conversational answers, which are streamed back to the frontend in real time so the analyst sees the response as it is generated.")
    
    doc.add_heading('Workflow Diagram', level=2)
    flow1 = ["User Input", "React Frontend", "FastAPI Backend", "Context Builder", "Groq LLM", "AI Generated Response"]
    for item in flow1:
        p = doc.add_paragraph(item)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if item != flow1[-1]:
            arrow = doc.add_paragraph('↓')
            arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Q2
    doc.add_heading('Q2. Prompt Chaining', level=1)
    doc.add_paragraph("Complex investigations in SentinelAI are handled through a chain of prompts rather than a single LLM call. An initial user request first triggers threat-context collection from the incident database. This context is then enriched with attack details drawn from correlation and honeypot logs. The enriched attack data is passed forward for MITRE ATT&CK mapping, which classifies the technique and tactic involved. The mapped result is combined with a Threat Intelligence lookup to add reputation and indicator data. Only after these stages does the system perform its core AI reasoning step, and the final response is generated using the accumulated context. At every stage, the output produced becomes the input consumed by the next stage, so accuracy compounds as the chain progresses.")
    
    doc.add_heading('Workflow Diagram', level=2)
    flow2 = ["Initial User Request", "Threat Context Collection", "Attack Enrichment", "MITRE ATT&CK Mapping", "Threat Intelligence Lookup", "AI Reasoning", "Final Response Generation"]
    for item in flow2:
        p = doc.add_paragraph(item)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if item != flow2[-1]:
            arrow = doc.add_paragraph('↓')
            arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Q3
    doc.add_heading('Q3. Agentic AI', level=1)
    doc.add_paragraph("The AI Security Copilot in SentinelAI behaves as an autonomous agent rather than a passive chatbot. On receiving a request, it first interprets user intent to decide whether the analyst wants a general conversation or a formal cybersecurity investigation. If investigation mode is selected, the agent independently retrieves related incidents, fetches supporting threat intelligence, performs MITRE ATT&CK mapping and correlates the activity against other known attacks — all without step-by-step instruction from the analyst. It then autonomously produces an investigation summary and a set of actionable recommendations, effectively planning and executing a multi-step security workflow on its own.")
    
    doc.add_heading('Workflow Diagram', level=2)
    flow3 = ["User Intent Understanding", "Mode Selection (Chat / Investigation)", "Incident Retrieval", "Threat Intelligence Fetch", "MITRE Mapping & Correlation", "Investigation Summary & Recommendations"]
    for item in flow3:
        p = doc.add_paragraph(item)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if item != flow3[-1]:
            arrow = doc.add_paragraph('↓')
            arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Q4
    doc.add_heading('Q4. RAG Based Question Answering', level=1)
    doc.add_paragraph("SentinelAI applies Retrieval-Augmented Generation so that every AI-generated answer is grounded in real security data rather than the model's general knowledge alone. Before a prompt reaches the Groq LLM, the system retrieves relevant records from the Incident Database, Threat Intelligence sources, Attack Logs, the MITRE ATT&CK Knowledge Base, past Security Reports, Honeypot Events and Correlated Incidents. The most relevant retrieved passages are injected directly into the LLM prompt as supporting context, allowing the model to produce accurate, evidence-backed answers about specific attacks and incidents rather than generic responses.")
    
    doc.add_heading('Workflow Diagram', level=2)
    flow4 = ["User Question", "Retrieval from Security Data Sources", "Context Injection", "Groq LLM", "Accurate, Grounded Response"]
    for item in flow4:
        p = doc.add_paragraph(item)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if item != flow4[-1]:
            arrow = doc.add_paragraph('↓')
            arrow.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Services Used
    doc.add_heading('Services Used', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Service'
    hdr_cells[1].text = 'Purpose'
    # Make header bold
    for cell in hdr_cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
    
    services = [
        ("FastAPI", "Backend API Framework"),
        ("React + Vite", "Frontend Application"),
        ("Groq API", "Large Language Model"),
        ("WebSocket", "Live Attack Feed"),
        ("SQLite / PostgreSQL", "Incident Storage"),
        ("Threat Intelligence Engine", "Threat Enrichment"),
        ("Correlation Engine", "Attack Correlation"),
        ("MITRE Mapping Module", "Threat Classification"),
        ("AI Security Copilot", "Investigation Assistant"),
        ("Docker", "Hybrid Deployment")
    ]
    
    for s, p in services:
        row_cells = table.add_row().cells
        row_cells[0].text = s
        row_cells[1].text = p

    # Technology Stack
    doc.add_heading('Technology Stack', level=1)
    techs = ['React.js', 'Vite', 'FastAPI', 'Python', 'Groq API', 'WebSockets', 'SQLite', 'PostgreSQL', 'Docker', 'Nginx', 'SQLAlchemy']
    for t in techs:
        doc.add_paragraph(t, style='List Bullet')

    # Conclusion
    doc.add_heading('Conclusion', level=1)
    doc.add_paragraph('SentinelAI successfully demonstrates all four Applied Agentic AI concepts through a single, cohesive cybersecurity platform. Its LLM workflow enables natural-language interaction with security data, its prompt chaining pipeline produces progressively refined investigation results, its agentic AI Security Copilot autonomously plans and executes multi-step analysis, and its Retrieval-Augmented Generation ensures every response remains grounded in real incident, threat-intelligence and honeypot data. Together, these capabilities allow SentinelAI to analyse attacks, assist analysts and generate AI-powered reports, meaningfully improving the speed and quality of security incident response.')

    # GitHub Repository
    doc.add_heading('GitHub Repository', level=1)
    p = doc.add_paragraph()
    runner = p.add_run('https://github.com/Vaishnav53/SentinelAI')
    runner.font.color.rgb = RGBColor(0, 0, 255)
    runner.font.underline = True

    doc.save('SentinelAI_Report.docx')

if __name__ == '__main__':
    create_report()
